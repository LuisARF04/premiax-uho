import csv
import io

from django.db import transaction

from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from auditoria.utils import registrar
from usuarios.permissions import IsAdmin
from .models import Area, Departamento
from .serializers import AreaSerializer, DepartamentoSerializer


class AreaViewSet(viewsets.ModelViewSet):
    queryset = Area.objects.select_related('responsable').prefetch_related('departamentos').all()
    serializer_class = AreaSerializer

    def get_permissions(self):
        if self.action in ('list', 'retrieve'):
            return [IsAuthenticated()]
        return [IsAdmin()]

    def perform_create(self, serializer):
        instance = serializer.save()
        registrar(self.request.user, 'Crear área', 'organizacion', instance.nombre, self.request)

    def perform_update(self, serializer):
        instance = serializer.save()
        registrar(self.request.user, 'Editar área', 'organizacion', instance.nombre, self.request)

    def destroy(self, request, *args, **kwargs):
        """RF-012: si el área tiene usuarios/departamentos asociados se exige confirmación explícita."""
        area = self.get_object()
        usuarios = area.usuarios.count()
        deptos = area.departamentos.count()
        if (usuarios or deptos) and request.query_params.get('confirmar') != 'true':
            return Response({
                'detail': f'El área "{area.nombre}" tiene {usuarios} usuario(s) y {deptos} departamento(s) asociados. '
                          'Si la elimina, los departamentos se borrarán y los usuarios quedarán sin área.',
                'requiere_confirmacion': True, 'usuarios': usuarios, 'departamentos': deptos,
            }, status=status.HTTP_409_CONFLICT)
        return super().destroy(request, *args, **kwargs)

    def perform_destroy(self, instance):
        registrar(self.request.user, 'Eliminar área', 'organizacion', instance.nombre, self.request)
        instance.delete()

    @action(detail=False, methods=['post'], permission_classes=[IsAdmin])
    def importar_excel(self, request):
        """
        RF-012: importa áreas y departamentos desde Excel (.xlsx) o CSV con columnas "area" y "departamento".
        La estructura y todas las filas se validan antes de importar; si hay errores no se importa nada.
        """
        archivo = request.FILES.get('archivo')
        if not archivo:
            return Response({'detail': 'Debe adjuntar un archivo.'}, status=status.HTTP_400_BAD_REQUEST)
        if archivo.size > 10 * 1024 * 1024:
            return Response({'detail': 'El archivo supera el tamaño máximo de 10 MB.'}, status=400)

        nombre = archivo.name.lower()
        try:
            if nombre.endswith('.xlsx'):
                encabezados, filas = _leer_xlsx(archivo)
            elif nombre.endswith('.csv'):
                encabezados, filas = _leer_csv(archivo)
            else:
                return Response({'detail': 'Formato no soportado. Use un archivo .xlsx o .csv.'}, status=400)
        except Exception:
            return Response({'detail': 'No se pudo leer el archivo. Verifique que no esté dañado y que sea .xlsx o .csv válido.'}, status=400)

        if 'area' not in encabezados:
            return Response({
                'detail': 'Estructura inválida: la primera fila debe tener las columnas "area" y "departamento".',
                'columnas_encontradas': encabezados,
            }, status=400)

        errores, validas = [], []
        for numero, fila in filas:
            nombre_area = fila.get('area', '')
            nombre_depto = fila.get('departamento', '')
            if not nombre_area:
                if nombre_depto:
                    errores.append(f'Fila {numero}: el departamento "{nombre_depto}" no tiene área.')
                continue
            if len(nombre_area) > 200 or len(nombre_depto) > 200:
                errores.append(f'Fila {numero}: los nombres no pueden superar 200 caracteres.')
                continue
            validas.append((nombre_area, nombre_depto))
        if errores:
            return Response({'detail': 'El archivo tiene errores; no se importó nada.', 'errores': errores[:50]}, status=400)
        if not validas:
            return Response({'detail': 'El archivo no contiene filas con datos.'}, status=400)

        creadas_areas, creados_deptos = 0, 0
        with transaction.atomic():
            for nombre_area, nombre_depto in validas:
                area, creada = Area.objects.get_or_create(nombre=nombre_area)
                creadas_areas += creada
                if nombre_depto:
                    _, creado = Departamento.objects.get_or_create(nombre=nombre_depto, area=area)
                    creados_deptos += creado
        registrar(request.user, 'Importar Excel/CSV', 'organizacion', f'{creadas_areas} áreas, {creados_deptos} departamentos', request)
        return Response({'areas_creadas': creadas_areas, 'departamentos_creados': creados_deptos, 'filas_procesadas': len(validas)})


def _normalizar(valor):
    return ('' if valor is None else str(valor)).strip()


def _leer_csv(archivo):
    lector = csv.DictReader(io.StringIO(archivo.read().decode('utf-8-sig')))
    encabezados = [_normalizar(c).lower() for c in (lector.fieldnames or [])]
    filas = []
    for numero, fila in enumerate(lector, start=2):
        filas.append((numero, {_normalizar(k).lower(): _normalizar(v) for k, v in fila.items() if k}))
    return encabezados, filas


def _leer_xlsx(archivo):
    from openpyxl import load_workbook
    hoja = load_workbook(archivo, read_only=True, data_only=True).active
    filas_crudas = hoja.iter_rows(values_only=True)
    encabezados = [_normalizar(c).lower() for c in next(filas_crudas, ())]
    filas = []
    for numero, valores in enumerate(filas_crudas, start=2):
        fila = {encabezados[i]: _normalizar(v) for i, v in enumerate(valores) if i < len(encabezados) and encabezados[i]}
        if any(fila.values()):
            filas.append((numero, fila))
    return encabezados, filas


class DepartamentoViewSet(viewsets.ModelViewSet):
    queryset = Departamento.objects.select_related('area', 'responsable').all()
    serializer_class = DepartamentoSerializer
    filterset_fields = ['area', 'activo']

    def get_permissions(self):
        if self.action in ('list', 'retrieve'):
            return [IsAuthenticated()]
        return [IsAdmin()]

    def perform_create(self, serializer):
        instance = serializer.save()
        registrar(self.request.user, 'Crear departamento', 'organizacion', instance.nombre, self.request)

    def perform_update(self, serializer):
        instance = serializer.save()
        registrar(self.request.user, 'Editar departamento', 'organizacion', instance.nombre, self.request)

    def destroy(self, request, *args, **kwargs):
        depto = self.get_object()
        usuarios = depto.usuarios.count()
        if usuarios and request.query_params.get('confirmar') != 'true':
            return Response({
                'detail': f'El departamento "{depto.nombre}" tiene {usuarios} usuario(s) asociados. '
                          'Si lo elimina, quedarán sin departamento.',
                'requiere_confirmacion': True, 'usuarios': usuarios,
            }, status=status.HTTP_409_CONFLICT)
        return super().destroy(request, *args, **kwargs)

    def perform_destroy(self, instance):
        registrar(self.request.user, 'Eliminar departamento', 'organizacion', instance.nombre, self.request)
        instance.delete()
